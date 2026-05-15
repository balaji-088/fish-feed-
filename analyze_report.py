from docx import Document
import os

def analyze_sample(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
    
    doc = Document(file_path)
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print("\n--- First 30 Paragraphs ---")
    for i, para in enumerate(doc.paragraphs[:30]):
        print(f"[{i}] Style: {para.style.name} | Text: {para.text[:100]}")

if __name__ == "__main__":
    analyze_sample("Mini Project report Sample.docx")
